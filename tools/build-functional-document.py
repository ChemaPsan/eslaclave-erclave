"""Build user manuals and QA guides from canonical Markdown with real Word tables."""
import argparse
import re
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def inline(paragraph, text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    for part in re.split(r'(\*\*.*?\*\*|`.*?`)', text):
        if not part:
            continue
        run = paragraph.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part)
        if part.startswith('**'):
            run.bold = True
        elif part.startswith('`'):
            run.font.name = 'Consolas'
            run.font.size = Pt(9)


def element(parent, tag, attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn(key), str(value))
    parent.append(node)
    return node


def table(document, rows):
    count = len(rows[0])
    grid = document.add_table(rows=0, cols=count)
    grid.autofit = False
    widths = [8.8, 3.9, 3.9] if count == 3 and rows[0][1:] == ['Local', 'QA'] else [5, 11.6] if count == 2 else [4.1, 3.5, 9] if count == 3 else [16.6/count]*count
    for col, width in zip(grid.columns, widths):
        col.width = Cm(width)
    borders = element(grid._tbl.tblPr, 'w:tblBorders', {})
    for edge in ['top','left','bottom','right','insideH','insideV']:
        element(borders, f'w:{edge}', {'w:val':'single','w:sz':4,'w:color':'D9D9D9'})
    for index, values in enumerate(rows):
        row = grid.add_row()
        element(row._tr.get_or_add_trPr(), 'w:cantSplit', {})
        if index == 0:
            element(row._tr.get_or_add_trPr(), 'w:tblHeader', {})
        for col, cell in enumerate(row.cells):
            cell.width = Cm(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = element(props,'w:tcMar',{})
            for side,size in [('top',85),('bottom',85),('left',100),('right',100)]:
                element(margins,f'w:{side}',{'w:w':size,'w:type':'dxa'})
            element(props,'w:shd',{'w:fill':'243746' if index == 0 else 'F2F5F7' if index%2 == 0 else 'FFFFFF'})
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1
            inline(p, values[col] if col<len(values) else '')
            for run in p.runs:
                run.font.size = Pt(9.5)
                if index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string('FFFFFF')
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build(source: Path, target: Path):
    content = source.read_text(encoding='utf-8')
    document = Document()
    # Some runtime templates inherit a blue paragraph rule under Title.
    # Manuals use typography and whitespace rather than that decorative rule.
    for border in document.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    section = document.sections[0]
    section.page_width,section.page_height = Cm(21.59),Cm(27.94)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2.495)
    normal = document.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.widow_control = True
    for name,size in [('Title',23),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
        style = document.styles[name]
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0,0,0)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(11)
        style.paragraph_format.space_after = Pt(6)
    for name in ['List Bullet','List Number']:
        document.styles[name].paragraph_format.space_after = Pt(4)
    footer = section.footer.paragraphs[0]
    footer.text = 'ERClave  |  Local / QA pendiente  |  14 septiembre 2026  |  '
    element(footer._p,'w:fldSimple',{'w:instr':'PAGE'})
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(70,70,70)
    lines = content.splitlines()
    i,code,title = 0,False,''
    while i<len(lines):
        line = lines[i].strip()
        i+=1
        if not line:
            continue
        if line.startswith('```'):
            code = not code
            continue
        if line.startswith('|') and not code:
            rows=[]
            while True:
                if not re.match(r'^\|[\s:|\-]+\|$',line):
                    rows.append([v.strip() for v in line.strip('|').split('|')])
                if i>=len(lines) or not lines[i].strip().startswith('|'):
                    break
                line=lines[i].strip()
                i+=1
            table(document,rows)
            continue
        match=re.match(r'^(#{1,4})\s+(.+)$',line)
        if match:
            level=len(match[1])
            heading=match[2].replace('**','').replace('`','')
            if level==1:
                title=heading
                document.add_paragraph(heading,style='Title')
            else:
                document.add_heading(heading,level=level-1)
        elif line=='---page---':
            document.add_page_break()
        elif line.startswith('- '):
            inline(document.add_paragraph(style='List Bullet'),line[2:])
        elif re.match(r'^\d+\.\s+',line):
            # Keep explicit numbering so unrelated procedures do not continue.
            p=document.add_paragraph()
            p.paragraph_format.left_indent=Cm(0.5)
            p.paragraph_format.first_line_indent=Cm(-0.5)
            inline(p,line)
        else:
            inline(document.add_paragraph(),line.removeprefix('> '))
    document.core_properties.title=title
    document.core_properties.author='ERClave'
    document.core_properties.subject='Manual operativo: base QA b63cdad y mejoras Local pendientes de promoción'
    target.parent.mkdir(parents=True,exist_ok=True)
    document.save(target)
    reopened=Document(target)
    assert reopened.paragraphs[0].style.name=='Title'
    assert len(reopened.tables)==sum(1 for idx,line in enumerate(lines) if line.startswith('|') and (idx==0 or not lines[idx-1].startswith('|')))
    print(f'{target.name}: {len(reopened.paragraphs)} paragraphs, {len(reopened.tables)} tables')


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('source',type=Path)
    parser.add_argument('target',type=Path)
    args=parser.parse_args()
    build(args.source,args.target)
